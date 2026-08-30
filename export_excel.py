
import sqlite3
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

DB = Path("data/market_monitor.db")
OUT = Path("output/투자_실시간수급_매매일지.xlsx")
SUMMARY = Path("output/summary.txt")


def export():
    OUT.parent.mkdir(exist_ok=True)

    con = sqlite3.connect(DB)
    rows = con.execute("""
        SELECT collected_at, code, name,
               foreign_sell, foreign_buy,
               foreign_net, foreign_net_change,
               total_sell, total_buy
        FROM foreign_member_flow
        ORDER BY collected_at
    """).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "국내수급"

    headers = [
        "수집시각", "종목코드", "종목명",
        "외국계매도합", "외국계매수합",
        "외국계순매수", "순매수변화",
        "전체매도합", "전체매수합"
    ]
    ws.append(headers)

    for r in rows:
        ws.append(r)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(horizontal="center")

    widths = [20, 12, 20, 16, 16, 16, 16, 16, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = w

    # 종목별 최신값 요약
    dash = wb.create_sheet("대시보드")
    dash.append(["종목코드", "종목명", "최신 외국계순매수", "최신 순매수변화"])

    latest = {}
    for r in rows:
        latest[r[1]] = r

    ranked = sorted(
        latest.values(),
        key=lambda x: x[5] if isinstance(x[5], (int, float)) else 0,
        reverse=True
    )

    for r in ranked:
        dash.append([r[1], r[2], r[5], r[6]])

    for cell in dash[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")

    for c in ["A","B","C","D"]:
        dash.column_dimensions[c].width = 22

    wb.save(OUT)

    lines = [
        f"NH PLUG 수급 모니터링 결과",
        f"생성시각: {datetime.now():%Y-%m-%d %H:%M:%S}",
        f"수집건수: {len(rows)}",
        "",
        "외국계 순매수 상위:"
    ]

    for r in ranked[:10]:
        lines.append(
            f"- {r[2] or r[1]}({r[1]}): "
            f"{r[5]:,}주 / 직전변화 {r[6]:,}주"
        )

    SUMMARY.write_text("\n".join(lines), encoding="utf-8")
    print(OUT)
    print(SUMMARY)


if __name__ == "__main__":
    export()
